"""Выгрузка участников в Excel."""
from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from .config import MEDIA_DIR
from .db import session_scope
from .models import Participant

PLATFORM_TITLE = {"tg": "Telegram", "vk": "ВКонтакте", "max": "МАКС"}

HEADERS = ["Номер", "ФИО", "Телефон", "Город", "Бонусная карта", "Площадка", "Дата регистрации"]


def export_participants(platform: Optional[str] = None) -> Path:
    """Создаёт .xlsx с участниками. platform=None — все площадки.

    Возвращает путь к файлу. Обязательные по ТЗ поля: телефон, ФИО, город.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Участники"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    with session_scope() as s:
        # только участники, загрузившие хотя бы один чек
        query = select(Participant).where(Participant.receipts.any())
        if platform:
            query = query.where(Participant.platform == platform)
        query = query.order_by(Participant.platform, Participant.number)
        participants = s.scalars(query).all()

        for i, p in enumerate(participants, start=2):
            ws.cell(row=i, column=1, value=p.display_number)
            ws.cell(row=i, column=2, value=p.full_name)
            ws.cell(row=i, column=3, value=p.phone)
            ws.cell(row=i, column=4, value=p.city)
            ws.cell(row=i, column=5, value=p.card_number)
            ws.cell(row=i, column=6, value=PLATFORM_TITLE.get(p.platform, p.platform))
            ws.cell(
                row=i,
                column=7,
                value=p.created_at.strftime("%d.%m.%Y %H:%M") if p.created_at else "",
            )

    # Ширина колонок
    widths = [10, 32, 18, 20, 20, 14, 20]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    suffix = platform or "all"
    path = MEDIA_DIR / f"participants_{suffix}_{stamp}.xlsx"
    wb.save(path)
    return path
